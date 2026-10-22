"""DOCX -> Markdown with python-docx. Word already stores headings, lists and reading order,
so no layout model or OCR is needed."""

from __future__ import annotations

import re
from importlib import metadata
from pathlib import Path
from typing import Any

from src.config import Settings
from src.providers.base import ParsedDocument

MAX_LABEL_CHARS = 60
MAX_HEADING_LEVEL = 3


class PythonDocxParser:
    name = "python_docx"

    def __init__(self, settings: Settings | None = None) -> None:
        self.settings = settings

    def parse(self, path: Path) -> ParsedDocument:
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        document = docx.Document(str(path))
        warnings: list[str] = []
        lines: list[str] = []

        # Designer templates often keep the name and contact details in the page header.
        header = [p.text.strip() for p in document.sections[0].header.paragraphs if p.text.strip()]
        lines += header

        for child in document.element.body.iterchildren():
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "p":
                lines += _paragraph(Paragraph(child, document))
            elif tag == "tbl":
                lines += _table(Table(child, document))

        if "txbxContent" in document.element.xml:
            warnings.append("document contains text boxes; their text may be missing")
        markdown = re.sub(r"\n{3,}", "\n\n", "\n\n".join(ln for ln in lines if ln)).strip()
        return ParsedDocument(
            markdown=markdown,
            n_pages=1,
            parser=self.name,
            parser_version=_version("python-docx"),
            warnings=warnings,
            pages=[markdown],
        )


def _paragraph(p: Any) -> list[str]:
    text = p.text.strip()
    if not text:
        return []
    style = (p.style.name if p.style is not None else "") or ""
    if style == "Title":
        return [f"# {text}"]
    heading = re.fullmatch(r"Heading (\d)", style)
    if heading:
        return ["#" * min(int(heading.group(1)) + 1, MAX_HEADING_LEVEL) + " " + text]
    ppr = p._p.pPr
    if style.startswith("List") or (ppr is not None and ppr.numPr is not None):
        return [f"- {text}"]
    runs = [r for r in p.runs if r.text.strip()]
    if runs and all(r.bold for r in runs) and len(text) <= MAX_LABEL_CHARS:
        return [f"**{text}**"]
    return [text]


def _table(table: Any) -> list[str]:
    """Emit each cell's paragraphs in row order, left to right. Word templates use tables as
    page columns, so keeping a cell's content together matters more than the grid."""
    out: list[str] = []
    for row in table.rows:
        seen = set()
        for cell in row.cells:
            if id(cell._tc) in seen:  # merged cells repeat
                continue
            seen.add(id(cell._tc))
            for p in cell.paragraphs:
                out += _paragraph(p)
    return out


def _version(package: str) -> str:
    try:
        return metadata.version(package)
    except metadata.PackageNotFoundError:
        return ""
