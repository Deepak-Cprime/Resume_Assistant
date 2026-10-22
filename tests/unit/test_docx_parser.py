from docx import Document

from src.providers.parsers.docx_parser import PythonDocxParser


def parse(tmp_path, build):
    doc = Document()
    build(doc)
    path = tmp_path / "t.docx"
    doc.save(path)
    return PythonDocxParser().parse(path)


def test_header_text_is_included(tmp_path):
    def build(d):
        d.sections[0].header.paragraphs[0].text = "Jane Doe jane@example.com"
        d.add_paragraph("Body text")

    md = parse(tmp_path, build).markdown
    assert md.startswith("Jane Doe jane@example.com")


def test_table_cells_are_read_left_to_right(tmp_path):
    def build(d):
        t = d.add_table(rows=1, cols=2)
        t.cell(0, 0).text = "SIDEBAR Python"
        t.cell(0, 1).text = "MAIN Senior Engineer"

    md = parse(tmp_path, build).markdown
    assert md.index("SIDEBAR") < md.index("MAIN")


def test_merged_cell_text_appears_once(tmp_path):
    def build(d):
        t = d.add_table(rows=1, cols=2)
        merged = t.cell(0, 0).merge(t.cell(0, 1))
        merged.text = "Merged heading"

    assert parse(tmp_path, build).markdown.count("Merged heading") == 1


def test_styles_map_to_markdown(tmp_path):
    def build(d):
        d.add_heading("Resume", 0)
        d.add_heading("Experience", 1)
        d.add_heading("Detail", 2)
        d.add_paragraph("Built things", style="List Bullet")
        d.add_paragraph().add_run("Staff Engineer, Acme").bold = True
        d.add_paragraph("Plain sentence about work.")

    md = parse(tmp_path, build).markdown
    assert md.splitlines()[0] == "# Resume"
    for expected in ("## Experience", "### Detail", "- Built things", "**Staff Engineer, Acme**"):
        assert expected in md
    assert "Plain sentence about work." in md


def test_empty_paragraphs_are_dropped(tmp_path):
    def build(d):
        d.add_paragraph("A")
        d.add_paragraph("   ")
        d.add_paragraph("B")

    assert parse(tmp_path, build).markdown == "A\n\nB"
