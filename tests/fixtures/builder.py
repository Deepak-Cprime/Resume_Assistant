"""Build synthetic resume files from the YAML specs in tests/fixtures/specs/.

Each spec is both the content and the ground truth (expected sections, jobs, YOE) for the
later milestones. All people, companies and contact details are fictional.
Layouts: single_column (PDF), two_column (PDF with a sidebar), scanned (rasterised PDF),
docx (Word document).
"""

from __future__ import annotations

import io
from datetime import datetime
from pathlib import Path
from typing import Any

import pymupdf
import yaml
from docx import Document

SPEC_DIR = Path(__file__).parent / "specs"
DEFAULT_HEADINGS = {
    "contact": "Contact",
    "summary": "Summary",
    "experience": "Experience",
    "education": "Education",
    "skills": "Skills",
    "certifications": "Certifications",
}
SINGLE_COLUMN_ORDER = [
    "name", "contact", "summary", "experience", "education", "skills", "certifications",
]
PAGE_W, PAGE_H, MARGIN = 612, 792, 54
# kind -> (font, size, line height, space before)
STYLES = {
    "name": ("hebo", 20, 26, 0),
    "heading": ("hebo", 13, 18, 12),
    "sub": ("hebo", 10.5, 15, 6),
    "body": ("helv", 10, 14, 0),
    "bullet": ("helv", 10, 14, 0),
    "small": ("heit", 9, 13, 0),
}
Block = tuple[str, str]


def load_specs() -> dict[str, dict[str, Any]]:
    return {
        spec["id"]: spec
        for spec in (yaml.safe_load(p.read_text("utf-8")) for p in sorted(SPEC_DIR.glob("*.yaml")))
    }


def fmt_date(spec: dict[str, Any], iso: str) -> str:
    if iso == "present":
        return spec.get("present_label", "Present")
    return datetime.strptime(iso, "%Y-%m").strftime(spec.get("date_format", "%b %Y"))


def _heading(spec: dict[str, Any], key: str) -> list[Block]:
    if spec.get("no_headings"):
        return []
    return [("heading", spec.get("headings", {}).get(key, DEFAULT_HEADINGS[key]))]


def section_blocks(spec: dict[str, Any], *, sidebar: bool = False) -> dict[str, list[Block]]:
    """Content blocks per section. In a sidebar, contact and skills are one item per line."""
    if sidebar:
        contact = [("small", spec["email"]), ("small", spec["phone"]), ("small", spec["location"])]
        skills = [("body", s) for s in spec["skills"]]
    else:
        contact = [("small", f"{spec['email']} | {spec['phone']} | {spec['location']}")]
        skills = [("body", ", ".join(spec["skills"]))]

    experience = _heading(spec, "experience")
    for job in spec["experience"]:
        dates = f"{fmt_date(spec, job['start'])} - {fmt_date(spec, job['end'])}"
        experience.append(("sub", f"{job['title']}, {job['company']}"))
        experience.append(("small", f"{dates} | {job['location']}"))
        experience += [("bullet", b) for b in job["bullets"]]

    education = _heading(spec, "education")
    for ed in spec["education"]:
        education.append(("sub", ed["degree"]))
        education.append(("small", f"{ed['school']}, {ed['year']}"))

    return {
        "name": [("name", spec["name"])],
        "contact": _heading(spec, "contact") + contact if sidebar else contact,
        "summary": _heading(spec, "summary") + [("body", spec["summary"])],
        "experience": experience,
        "education": education,
        "skills": _heading(spec, "skills") + skills,
        "certifications": _heading(spec, "certifications")
        + [("bullet", c) for c in spec["certifications"]],
    }


def _wrap(text: str, font: str, size: float, width: float) -> list[str]:
    lines: list[str] = []
    current = ""
    for word in text.split():
        trial = f"{current} {word}".strip()
        if current and pymupdf.get_text_length(trial, fontname=font, fontsize=size) > width:
            lines.append(current)
            current = word
        else:
            current = trial
    return [*lines, current]


class _Column:
    """A text column that flows onto new pages when it runs out of room."""

    def __init__(self, doc: pymupdf.Document, x0: float, x1: float) -> None:
        self.doc, self.x0, self.x1 = doc, x0, x1
        self.page_idx, self.y = 0, float(MARGIN)
        self._ensure_page()

    def _ensure_page(self) -> None:
        while len(self.doc) <= self.page_idx:
            self.doc.new_page(width=PAGE_W, height=PAGE_H)

    def write(self, blocks: list[Block]) -> None:
        for kind, text in blocks:
            font, size, line_h, before = STYLES[kind]
            indent = 12 if kind == "bullet" else 0
            if self.y > MARGIN:
                self.y += before
            for i, line in enumerate(_wrap(text, font, size, self.x1 - self.x0 - indent)):
                if self.y + line_h > PAGE_H - MARGIN:
                    self.page_idx += 1
                    self.y = float(MARGIN)
                    self._ensure_page()
                page = self.doc[self.page_idx]
                self.y += line_h
                if kind == "bullet" and i == 0:
                    page.insert_text((self.x0, self.y), "•", fontname="helv", fontsize=size)
                page.insert_text((self.x0 + indent, self.y), line, fontname=font, fontsize=size)


def build_pdf(spec: dict[str, Any]) -> bytes:
    doc = pymupdf.open()
    if spec["layout"] == "two_column":
        side_w, gap = spec.get("sidebar_width", 150), 24
        blocks = section_blocks(spec, sidebar=True)
        side_blocks = (
            blocks["contact"] + blocks["skills"] + blocks["certifications"] + blocks["education"]
        )
        main_blocks = blocks["name"] + blocks["summary"] + blocks["experience"]
        right_edge = PAGE_W - MARGIN
        side_left = (MARGIN, MARGIN + side_w)
        main_right = (MARGIN + side_w + gap, right_edge)
        main_left = (MARGIN, right_edge - side_w - gap)
        side_right = (right_edge - side_w, right_edge)
        # Write left-to-right, as the content stream of a real two-column PDF usually is.
        if spec.get("sidebar", "left") == "left":
            _Column(doc, *side_left).write(side_blocks)
            _Column(doc, *main_right).write(main_blocks)
        else:
            _Column(doc, *main_left).write(main_blocks)
            _Column(doc, *side_right).write(side_blocks)
    else:
        blocks = section_blocks(spec)
        column = _Column(doc, MARGIN, PAGE_W - MARGIN)
        for key in SINGLE_COLUMN_ORDER:
            column.write(blocks[key])
    data = doc.tobytes()
    doc.close()
    return _rasterise(data) if spec["layout"] == "scanned" else data


def _rasterise(pdf_bytes: bytes, dpi: int = 150) -> bytes:
    """Turn a digital PDF into an image-only one, like a scan (no text layer)."""
    src = pymupdf.open(stream=pdf_bytes, filetype="pdf")
    out = pymupdf.open()
    for page in src:
        pix = page.get_pixmap(dpi=dpi)
        new = out.new_page(width=page.rect.width, height=page.rect.height)
        new.insert_image(new.rect, pixmap=pix)
    data = out.tobytes()
    src.close()
    out.close()
    return data


def build_docx(spec: dict[str, Any]) -> bytes:
    doc = Document()
    blocks = section_blocks(spec)
    for key in SINGLE_COLUMN_ORDER:
        for kind, text in blocks[key]:
            if kind == "name":
                doc.add_heading(text, 0)
            elif kind == "heading":
                doc.add_heading(text, 1)
            elif kind == "sub":
                doc.add_paragraph().add_run(text).bold = True
            elif kind == "bullet":
                doc.add_paragraph(text, style="List Bullet")
            else:
                doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build(spec: dict[str, Any]) -> tuple[bytes, str]:
    """Return (file bytes, extension) for one spec."""
    if spec["layout"] == "docx":
        return build_docx(spec), ".docx"
    return build_pdf(spec), ".pdf"


def write_all(out_dir: Path) -> dict[str, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    paths: dict[str, Path] = {}
    for spec_id, spec in load_specs().items():
        data, ext = build(spec)
        paths[spec_id] = out_dir / f"{spec_id}{ext}"
        paths[spec_id].write_bytes(data)
    return paths
